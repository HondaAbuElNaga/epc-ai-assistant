"""Builds a tiny fake Ghent DSLIB folder for tests.

Same layout as the real DSLIB v3.4 files (specs/2026-10-08-ghent-project-db/format_notes.md),
but every name and number is made up: the real data has no license, so none of it is copied.

Three projects, one per group, each with a quirk seen in the real data:
- C2011-05 (group A): WBS summary row, durations `3d` / `1d 2h` / `4h`, Tracking Overview.
- C2023-01 (group B): TP sheet with a trailing space, no Tracking Overview, text IDs,
  duplicate activity ID, `58 days`, no `Baseline End` / `Total Cost` columns, BAC `€743 676`.
- C2012-03 (group C): empty Tracking Overview, `Sussessors` typo, BAC `-`.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string

GREEN, YELLOW, ORANGE = "FF00FF00", "FFFFFF00", "FFFF8000"

BASELINE_HEADER = [
    "ID", "Name", "WBS", "Predecessors", "Successors", "Baseline Start", "Baseline End",
    "Duration", "Resource Demand", "Resource Cost", "Fixed Cost", "Cost/Hour", "Variable Cost",
    "Total Cost", None, None, "Baseline duration (in calendar days)",
]  # fmt: skip

TRACKING_HEADER = [
    "Name", "Start Tracking Period", "Status date", "Planned Value (PV)", "Earned Value (EV)",
    "Actual Cost (AC)", "Earned Schedule (ES)", "Schedule Variance (SV)",
    "Schedule Performance Index (SPI)", "Cost Variance (CV)", "Cost Performance Index (CPI)",
    "Schedule Variance (SV(t))", "Schedule Performance Index (SPI(t))", "p-factor",
    "EAC(t)-PV (PF=1)",
]  # fmt: skip

TP_HEADER = [
    "ID", "Name", "Baseline Start", "Baseline End", "Duration", "Resource Demand",
    "Resource Cost", "Fixed Cost", "Cost/Hour", "Variable Cost", "Total Cost", "Actual Start",
    "Actual Duration", "PAC", "PRC", "Remaining Duration", "PAC Dev", "PRC Dev", "Actual Cost",
    "Remaining Cost", "Percentage Completed", "Tracking", "Earned Value (EV)",
    "Planned Value (PV)",
]  # fmt: skip

D = datetime


def _sheet(wb: Workbook, title: str, rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for row in rows:
        ws.append(row)
    return ws


# --- per-project workbooks ------------------------------------------------------------------


def _project_a(path: Path, drop: str | None) -> None:
    header = [None if h == drop else h for h in BASELINE_HEADER]
    wb = Workbook()
    _sheet(wb, "Baseline Schedule", first=True, rows=[
        ["General", None, None, "Relations", None, "Baseline", None, None, "Resource Demand",
         None, "Baseline Costs"],
        header,
        [0, "Test Telecom", "1", None, None, D(2011, 5, 2, 8), D(2011, 5, 13, 17), "10d",
         None, None, 300.0, None, None, 1300.0, None, None, 11.375],
        [1, "Design", "1.1", None, "FS3", D(2011, 5, 2, 8), D(2011, 5, 4, 17), "3d",
         None, None, 0, 0, 0, 1000.0, None, None, 2.375],
        [2, "Order parts", "1.1.1", None, "FS3", D(2011, 5, 2, 8), D(2011, 5, 4, 17), "3d",
         "engineer", 600.0, 0, 0, 0, 600.0, None, None, 2.375],
        [4, "Approve design", "1.1.2", None, "FS3", D(2011, 5, 2, 8), D(2011, 5, 2, 12), "4h",
         "engineer", 400.0, 0, 0, 0, 400.0, None, None, 0.167],
        [3, "Install", "1.2", "2FS;4FS", None, D(2011, 5, 5, 8), D(2011, 5, 6, 10), "1d 2h",
         None, None, 300.0, 0, 0, 300.0, None, None, 1.083],
    ])  # fmt: skip
    _sheet(wb, "Resources", [["General"], ["ID", "Name", "Type"], [1, "engineer", "Renewable"]])
    _sheet(wb, "Project Control - TP1", [["", "TP Status Date", D(2011, 5, 6, 17)]])
    _sheet(wb, "TP2", [["", "TP Status Date", D(2011, 5, 16, 17)]])
    _sheet(wb, "Agenda", [["Working Hours", None, None, "Working Days"]])
    _sheet(wb, "Tracking Overview", rows=[
        ["General", None, None, "EVM Performance Measures", None, None, None, "EVM Forecasting"],
        TRACKING_HEADER,
        ["06/05, 2011", D(2011, 5, 2, 8), D(2011, 5, 6, 17), 800.0, 700.0, 750.0,
         D(2011, 5, 5, 12), -100.0, 0.875, -50.0, 0.9333, "-1d 4h", 0.9, 0.95, D(2011, 5, 16)],
        ["Actual Schedule", D(2011, 5, 6, 17), D(2011, 5, 16, 17), 1300.0, 1300.0, 1400.0,
         D(2011, 5, 13, 17), 0.0, 1.0, -100.0, 0.9286, "-1d", 0.95, 1.0, D(2011, 5, 16)],
    ])  # fmt: skip
    wb.save(path)


def _project_b(path: Path) -> None:
    # Newer layout: no Baseline End / Total Cost / calendar-days columns, IDs stored as text
    header = [
        "ID", "Name", "WBS", "Predecessors", "Successors", "Baseline Start", None, "Duration",
        "Resource Demand", None, "Fixed Cost", None, "Variable Cost",
    ]  # fmt: skip
    wb = Workbook()
    _sheet(wb, "Baseline Schedule", first=True, rows=[
        ["General", None, None, "Relations", None, "Baseline"],
        header,
        [0, "Test Renovation"],
        ["2", "Scaffolding", None, None, "FS3", D(2022, 1, 17, 8), None, "58 days",
         "worker[3,00 #8]", None, 5000.0, None, 0],
        ["3", "Demolition", None, "2FS", None, D(2022, 4, 4, 8), None, "5d", None, None,
         2000.0, None, 0],
        ["3", "Demolition (copy)", None, None, None, D(2022, 4, 4, 8), None, "0", None, None,
         0.0, None, 0],
    ])  # fmt: skip
    _sheet(wb, "Resources", [["General"], ["ID", "Name", "Type"], [1, "worker", "Renewable"]])
    _sheet(wb, "Risk Analysis", [["General"]])
    _sheet(wb, "Project Control - TP1 ", rows=[
        [None, "TP Status Date", D(2022, 2, 1, 17), None, "TP Name", "check 1"],
        [],
        ["General", None, "Baseline", None, None, "Resource Demand", None, "Baseline Costs",
         None, None, None, "Tracking"],
        TP_HEADER,
        [0, "Test Renovation"],
        ["2", "Scaffolding", None, None, None, None, None, None, None, None, None,
         D(2022, 1, 17, 8), "10d", None, None, None, None, None, 900.0, None, 0.2],
    ])  # fmt: skip
    _sheet(wb, "Agenda", [["Working Hours"]])
    wb.save(path)


def _project_c(path: Path) -> None:
    header = [
        "ID", "Name", "WBS", "Predecessors", "Sussessors", "Baseline Start", "Baseline End",
        "Duration", "Resource Demand", "Resource Cost", "Fixed Cost", "Cost/Hour",
        "Variable Cost", "Total Cost",
    ]  # fmt: skip
    wb = Workbook()
    _sheet(wb, "Baseline Schedule", first=True, rows=[
        ["General"],
        header,
        [0, "Test Day Care", "1", None, None, D(2012, 3, 1, 8), D(2012, 3, 2, 17), "2d",
         None, None, 0, None, None, 500.0],
        [1, "Plan", "1.1", None, "FS2", D(2012, 3, 1, 8), D(2012, 3, 1, 17), "1d", None,
         None, 200.0, 0, 0, 200.0],
        [2, "Build", "1.2", "1FS", None, D(2012, 3, 2, 8), D(2012, 3, 2, 17), "1d", None,
         None, 300.0, 0, 0, 300.0],
    ])  # fmt: skip
    _sheet(wb, "Resources", [["General"]])
    _sheet(wb, "Agenda", [["Working Hours"]])
    _sheet(wb, "Tracking Overview", [["General"], TRACKING_HEADER])  # header, no periods
    wb.save(path)


# --- analysis sheet -------------------------------------------------------------------------

# One dict per project: DSLIB column letter -> value (fills handled separately)
SUMMARY = [
    {"A": "C2011-05", "B": "Test Telecom", "C": "Tester One", "D": "IT", "E": "Telecom",
     "K": 4, "L": 10, "M": 1300, "N": "Y", "R": 0.1, "S": -0.0769, "T": 0.5, "U": 0.4,
     "V": 0.3, "W": 0.2, "X": 0.8, "Y": "Irregular", "CE": 11, "CF": 1400},
    {"A": "C2023-01", "B": "Test Renovation", "C": "Tester Two",
     "D": "Construction (Civil)", "E": None, "K": 3, "L": 63, "M": "€743 676", "N": "N",
     "R": "-", "S": "-", "T": 0.1, "U": 0.2, "V": 0.3, "W": 0.4, "X": "N/A", "Y": "N/A",
     "CE": "N/A", "CF": "N/A"},
    {"A": "C2012-03", "B": "Test Day Care", "C": "Tester Three",
     "D": "Construction (civil)", "E": "Child care", "K": 2, "L": 2, "M": "-", "N": "N/A",
     "R": "-", "S": "-", "T": 1, "U": 0, "V": 1, "W": 0, "X": "-", "Y": "regular",
     "CE": "N/A", "CF": "€\xa0464 186,97"},
]  # fmt: skip

FILLS = [  # (F baseline, G risk, H control) completeness colours
    (GREEN, YELLOW, GREEN),
    (YELLOW, YELLOW, "FFFF8001"),
    (ORANGE, GREEN, None),
]


def _analysis_sheet(path: Path) -> None:
    wb = Workbook()
    _sheet(wb, "Overview", [["Dynamic Scheduling Library (DSLIB): test copy"]], first=True)
    ws = wb.create_sheet("DSLIB")
    groups = {"A": "Project database", "F": "Completeness", "K": "General", "R": "Tracking",
              "T": "Network topology", "CE": "Real"}  # fmt: skip
    names = {"A": "Code", "B": "Project name", "C": "Submitted by", "D": "Sector",
             "E": "Keywords", "F": "Baseline Schedule", "G": "Risk Analysis",
             "H": "Project Control", "K": "# activities", "L": "PD (days)", "M": "BAC",
             "N": "Resources", "R": "Early/late", "S": "Under/over budget", "T": "SP",
             "U": "AD", "V": "LA", "W": "TF", "X": "RI", "Y": "Regularity",
             "CE": "Duration", "CF": "Cost"}  # fmt: skip
    for col, text in groups.items():
        ws.cell(1, column_index_from_string(col), text)
    for col, text in names.items():
        ws.cell(3, column_index_from_string(col), text)
    for i, (values, fills) in enumerate(zip(SUMMARY, FILLS, strict=True)):
        row = 4 + i
        for col, value in values.items():
            ws.cell(row, column_index_from_string(col), value)
        for col, rgb in zip("FGH", fills, strict=True):
            if rgb:
                ws[f"{col}{row}"].fill = PatternFill("solid", fgColor=rgb)
    # Summary statistics below the projects (not projects)
    ws.cell(9, 3, "SUMMARY STATISTICS")
    ws.cell(10, 3, "Construction (civil)")
    ws.cell(10, 5, 2)
    _sheet(wb, "Updates", [["Details about the updated files"]])
    _sheet(wb, "Known Problems", [["Known problems"], ["Project ID", "Update Type", "Comments"]])
    wb.save(path)


def build_dslib(root: Path, *, drop_baseline_column: str | None = None) -> Path:
    """Write the fake DSLIB folder under `root` and return its path.

    `drop_baseline_column` blanks one header cell of project C2011-05's Baseline Schedule, to
    test the missing-column error.
    """
    dslib = root / "DSLIB 3.4"
    excel = dslib / "Excel"
    excel.mkdir(parents=True)
    _analysis_sheet(dslib / "DSLIB_Analysis_Sheet.xlsx")
    _project_a(excel / "C2011-05 Test Telecom.xlsx", drop_baseline_column)
    _project_b(excel / "C2023-01 Test Renovation.xlsx")
    _project_c(excel / "C2012-3 Test Day Care.xlsx")  # unpadded ID, as in some real names
    return dslib
